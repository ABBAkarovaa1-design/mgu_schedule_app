import os, sqlite3, secrets, shutil
from datetime import datetime
from pathlib import Path
from fastapi import FastAPI, Request, Form, UploadFile, File
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
from openpyxl import load_workbook

BASE = Path(__file__).parent
DATA = BASE / "data"
DB = DATA / "schedule.db"
STATIC = BASE / "static"
DATA.mkdir(exist_ok=True)
STATIC.mkdir(exist_ok=True)

app = FastAPI(title="МГУ Расписание")
app.add_middleware(SessionMiddleware, secret_key=os.getenv("SESSION_SECRET", secrets.token_hex(32)))
app.mount("/static", StaticFiles(directory=STATIC), name="static")

DAYS = ["Понедельник","Вторник","Среда","Четверг","Пятница","Суббота","Воскресенье"]

def conn():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c

def init_db():
    c = conn()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS groups(
      id INTEGER PRIMARY KEY, name TEXT UNIQUE NOT NULL, faculty TEXT DEFAULT ''
    );
    CREATE TABLE IF NOT EXISTS lessons(
      id INTEGER PRIMARY KEY,
      group_id INTEGER NOT NULL,
      week INTEGER DEFAULT 0,
      day TEXT NOT NULL,
      start TEXT NOT NULL,
      end TEXT NOT NULL,
      subject TEXT NOT NULL,
      teacher TEXT DEFAULT '',
      room TEXT DEFAULT '',
      lesson_type TEXT DEFAULT '',
      note TEXT DEFAULT '',
      FOREIGN KEY(group_id) REFERENCES groups(id) ON DELETE CASCADE
    );
    CREATE INDEX IF NOT EXISTS idx_lessons_group ON lessons(group_id, week, day);
    """)
    c.commit(); c.close()

init_db()

def admin_ok(request: Request):
    return request.session.get("admin") is True

def admin_page():
    return (STATIC/"admin.html").read_text(encoding="utf-8")

@app.get("/", response_class=HTMLResponse)
async def home():
    return (STATIC/"index.html").read_text(encoding="utf-8")

@app.get("/api/groups")
async def groups():
    c=conn(); rows=c.execute("SELECT id,name,faculty FROM groups ORDER BY name").fetchall(); c.close()
    return [dict(x) for x in rows]

@app.get("/api/schedule/{group_id}")
async def schedule(group_id:int, week:int=0):
    c=conn()
    rows=c.execute("""SELECT id,day,start,end,subject,teacher,room,lesson_type,note
                      FROM lessons WHERE group_id=? AND week=? ORDER BY day,start""",(group_id,week)).fetchall()
    c.close()
    out={d:[] for d in DAYS}
    for r in rows: out[r["day"]].append(dict(r))
    return out

@app.get("/admin", response_class=HTMLResponse)
async def admin(request: Request):
    if not admin_ok(request):
        return RedirectResponse("/admin/login")
    return admin_page()

@app.get("/admin/login", response_class=HTMLResponse)
async def login_page():
    return """<!doctype html><html><body style="font-family:system-ui;background:#151218;color:white;padding:40px">
    <h2>МГУ — вход в админ-панель</h2><form method="post"><input name="password" type="password" placeholder="Пароль" style="padding:12px">
    <button style="padding:12px">Войти</button></form></body></html>"""

@app.post("/admin/login")
async def login(request:Request, password:str=Form(...)):
    if secrets.compare_digest(password, os.getenv("ADMIN_PASSWORD","change-this-password")):
        request.session["admin"]=True
        return RedirectResponse("/admin", status_code=303)
    return HTMLResponse("Неверный пароль", status_code=401)

@app.post("/admin/logout")
async def logout(request:Request):
    request.session.clear()
    return RedirectResponse("/")

@app.get("/api/admin/groups")
async def admin_groups(request:Request):
    if not admin_ok(request): return JSONResponse({"error":"auth"},401)
    return await groups()

@app.post("/api/admin/groups")
async def add_group(request:Request):
    if not admin_ok(request): return JSONResponse({"error":"auth"},401)
    data=await request.json(); c=conn()
    try:
        c.execute("INSERT INTO groups(name,faculty) VALUES(?,?)",(data["name"],data.get("faculty",""))); c.commit()
    except sqlite3.IntegrityError: return JSONResponse({"error":"Группа уже существует"},400)
    finally: c.close()
    return {"ok":True}

@app.delete("/api/admin/groups/{gid}")
async def del_group(gid:int, request:Request):
    if not admin_ok(request): return JSONResponse({"error":"auth"},401)
    c=conn(); c.execute("DELETE FROM lessons WHERE group_id=?",(gid,)); c.execute("DELETE FROM groups WHERE id=?",(gid,)); c.commit(); c.close()
    return {"ok":True}

@app.post("/api/admin/lessons")
async def add_lesson(request:Request):
    if not admin_ok(request): return JSONResponse({"error":"auth"},401)
    d=await request.json()
    c=conn(); c.execute("""INSERT INTO lessons(group_id,week,day,start,end,subject,teacher,room,lesson_type,note)
    VALUES(?,?,?,?,?,?,?,?,?,?)""",(d["group_id"],d.get("week",0),d["day"],d["start"],d["end"],d["subject"],d.get("teacher",""),d.get("room",""),d.get("lesson_type",""),d.get("note","")))
    c.commit(); c.close(); return {"ok":True}

@app.put("/api/admin/lessons/{lid}")
async def update_lesson(lid:int, request:Request):
    if not admin_ok(request): return JSONResponse({"error":"auth"},401)
    d=await request.json(); c=conn()
    c.execute("""UPDATE lessons SET week=?,day=?,start=?,end=?,subject=?,teacher=?,room=?,lesson_type=?,note=? WHERE id=?""",
              (d.get("week",0),d["day"],d["start"],d["end"],d["subject"],d.get("teacher",""),d.get("room",""),d.get("lesson_type",""),d.get("note",""),lid))
    c.commit(); c.close(); return {"ok":True}

@app.delete("/api/admin/lessons/{lid}")
async def del_lesson(lid:int, request:Request):
    if not admin_ok(request): return JSONResponse({"error":"auth"},401)
    c=conn(); c.execute("DELETE FROM lessons WHERE id=?",(lid,)); c.commit(); c.close(); return {"ok":True}

@app.get("/api/admin/all-lessons")
async def all_lessons(request:Request, group_id:int):
    if not admin_ok(request): return JSONResponse({"error":"auth"},401)
    c=conn(); rows=c.execute("SELECT * FROM lessons WHERE group_id=? ORDER BY week,day,start",(group_id,)).fetchall(); c.close()
    return [dict(x) for x in rows]

@app.post("/api/admin/import")
async def import_excel(request:Request, file:UploadFile=File(...)):
    if not admin_ok(request): return JSONResponse({"error":"auth"},401)
    tmp=DATA/"upload.xlsx"; tmp.write_bytes(await file.read())
    wb=load_workbook(tmp, data_only=True)
    ws=wb.active
    headers=[str(x.value).strip().lower() if x.value is not None else "" for x in ws[1]]
    required=["group","day","start","end","subject"]
    if not all(x in headers for x in required):
        return JSONResponse({"error":"Нужны колонки: group, day, start, end, subject; дополнительно teacher, room, faculty, week, lesson_type, note"},400)
    idx={h:i for i,h in enumerate(headers)}
    c=conn(); imported=0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[idx["group"]] or not row[idx["subject"]]: continue
        g=str(row[idx["group"]]).strip(); faculty=str(row[idx.get("faculty",-1)] or "") if idx.get("faculty",-1)>=0 else ""
        c.execute("INSERT OR IGNORE INTO groups(name,faculty) VALUES(?,?)",(g,faculty))
        gid=c.execute("SELECT id FROM groups WHERE name=?",(g,)).fetchone()["id"]
        def val(k,default=""): return str(row[idx[k]] if k in idx and row[idx[k]] is not None else default).strip()
        week=int(row[idx["week"]] or 0) if "week" in idx and row[idx["week"]] is not None else 0
        c.execute("""INSERT INTO lessons(group_id,week,day,start,end,subject,teacher,room,lesson_type,note)
        VALUES(?,?,?,?,?,?,?,?,?,?)""",(gid,week,val("day"),val("start"),val("end"),val("subject"),val("teacher"),val("room"),val("lesson_type"),val("note")))
        imported+=1
    c.commit(); c.close(); tmp.unlink(missing_ok=True)
    return {"ok":True,"imported":imported}

@app.post("/api/admin/backup")
async def backup(request:Request):
    if not admin_ok(request): return JSONResponse({"error":"auth"},401)
    stamp=datetime.now().strftime("%Y%m%d_%H%M%S"); target=DATA/f"schedule_backup_{stamp}.db"; shutil.copy2(DB,target)
    return {"ok":True,"file":target.name}
